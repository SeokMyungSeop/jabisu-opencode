#!/usr/bin/env python3
"""CSV 파일을 보기 좋은 엑셀(.xlsx) 통합 문서로 바꾼다.

사용법:
    python csv2xlsx.py 출력.xlsx 표1.csv [표2.csv ...] [--names 예산,일정] [--font "맑은 고딕"] [--size 10]

CSV 하나가 시트 하나가 된다(시트 이름은 --names, 없으면 파일 이름). 첫 줄은 머리글이 되어 굵게·배경색·
틀 고정·자동 필터가 들어간다. 숫자(1,500,000 같은 쉼표 숫자 포함)·백분율·날짜(2026-10-20, 2026.10.20)는
값으로 바꾸고, =로 시작하는 칸은 수식으로 둔다. 0으로 시작하는 숫자(우편번호, 전화번호)는 글자로 둔다.
열 너비는 내용에 맞춘다. CSV는 UTF-8(BOM 포함)이나 CP949(EUC-KR)를 읽는다.
"""

# 자비수: 문서 스킬이 함께 쓰는 Python 패키지 폴더를 먼저 찾는다(scripts/jabisu_python.py).
import os as _os
import sys as _sys
_sys.path.insert(0, _os.path.dirname(_os.path.abspath(__file__)))
import jabisu_python  # noqa: E402,F401

import argparse
import csv
import datetime
import io
import os
import re
import unicodedata

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NUMBER = re.compile(r"^[-+]?(\d{1,3}(,\d{3})+|\d+)(\.\d+)?$")
PERCENT = re.compile(r"^[-+]?\d+(\.\d+)?%$")
DATE = re.compile(r"^(\d{4})[-./](\d{1,2})[-./](\d{1,2})\.?$")
INVALID_SHEET = re.compile(r"[\[\]:*?/\\]")


def main():
    parser = argparse.ArgumentParser(description="CSV 파일을 엑셀(.xlsx) 통합 문서로 바꾼다.")
    parser.add_argument("output", help="만들 .xlsx 파일")
    parser.add_argument("inputs", nargs="+", help="CSV 파일 (하나가 시트 하나)")
    parser.add_argument("--names", help="시트 이름들, 쉼표로 구분")
    parser.add_argument("--font", default="맑은 고딕", help="글꼴 (기본: 맑은 고딕)")
    parser.add_argument("--size", type=float, default=10, help="글자 크기 pt (기본: 10)")
    args = parser.parse_args()

    names = [name.strip() for name in args.names.split(",")] if args.names else []
    workbook = Workbook()
    workbook.remove(workbook.active)
    for index, file in enumerate(args.inputs):
        name = names[index] if index < len(names) else os.path.splitext(os.path.basename(file))[0]
        sheet = workbook.create_sheet(unique(workbook, INVALID_SHEET.sub("_", name)[:31] or f"Sheet{index + 1}"))
        fill(sheet, read(file), args.font, args.size)
    workbook.save(args.output)
    print(f"만들었습니다: {args.output} (시트 {len(workbook.sheetnames)}개: {', '.join(workbook.sheetnames)})")


def read(file):
    data = open(file, "rb").read()
    for encoding in ("utf-8-sig", "cp949"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise SystemExit(f"{file}: UTF-8도 CP949도 아닙니다. 인코딩을 확인하세요.")
    return [row for row in csv.reader(io.StringIO(text))]


def unique(workbook, name):
    candidate, number = name, 2
    while candidate in workbook.sheetnames:
        suffix = f" ({number})"
        candidate, number = name[: 31 - len(suffix)] + suffix, number + 1
    return candidate


def fill(sheet, rows, font, size):
    border = Side(style="thin", color="BFBFBF")
    widths = {}
    for row_index, row in enumerate(rows, start=1):
        for column, text in enumerate(row, start=1):
            value, number_format = convert(text) if row_index > 1 else (text, None)
            cell = sheet.cell(row=row_index, column=column, value=value)
            cell.font = Font(name=font, size=size, bold=row_index == 1)
            cell.border = Border(left=border, right=border, top=border, bottom=border)
            if number_format:
                cell.number_format = number_format
            if row_index == 1:
                cell.fill = PatternFill("solid", fgColor="D9E1F2")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            widths[column] = max(widths.get(column, 0), display_width(text))
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = min(max(width + 2, 8), 60)
    if rows:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions


def convert(text):
    """칸의 글을 엑셀 값과 표시 형식으로 바꾼다."""
    value = text.strip()
    if value == "":
        return None, None
    if value.startswith("="):
        return value, None
    if PERCENT.match(value):
        number = float(value[:-1]) / 100
        return number, "0%" if "." not in value else "0.0%"
    plain = value.replace(",", "")
    if NUMBER.match(value) and not re.match(r"^[-+]?0\d", plain):
        if "." in plain:
            return float(plain), "#,##0.00" if "," in value else None
        # 쉼표를 쓴 숫자만 천 단위 쉼표로 보인다. 연도나 번호(20261020)는 그대로 둔다.
        return int(plain), "#,##0" if "," in value else None
    date = DATE.match(value)
    if date:
        try:
            return datetime.date(int(date.group(1)), int(date.group(2)), int(date.group(3))), "yyyy-mm-dd"
        except ValueError:
            return value, None
    return value, None


def display_width(text):
    """한글처럼 넓은 글자는 두 칸으로 센다."""
    return sum(2 if unicodedata.east_asian_width(char) in "WF" else 1 for char in str(text))


if __name__ == "__main__":
    main()
